from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_rtl(paragraph) -> None:
    """Set Word's bidi flag; alignment alone is not enough for Arabic."""
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)


def write_text_file(text: str, output_path: str | Path) -> Path:
    """Write OCR text as UTF-8 and return the created path."""
    path = Path(output_path)
    path.write_text(text, encoding="utf-8")
    return path


def write_docx_file(text: str, output_path: str | Path) -> Path:
    """Create a readable right-to-left Arabic OCR document."""
    path = Path(output_path)
    document = Document()
    document.add_heading("Arabic Handwritten Document - OCR", level=1)
    for paragraph_text in text.splitlines() or [""]:
        paragraph = document.add_paragraph(paragraph_text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl(paragraph)
        paragraph.paragraph_format.space_after = 0
    document.save(path)
    return path
